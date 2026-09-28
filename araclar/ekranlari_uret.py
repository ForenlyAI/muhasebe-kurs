"""Derslerin ekran görüntülerini baştan üretir (gerçek .xlsx → görüntüleyici, gerçek komut çıktısı).

    uv run --project araclar python araclar/ekranlari_uret.py            # hepsi
    uv run --project araclar python araclar/ekranlari_uret.py 2.2 3.3    # yalnız bu dersler

Önkoşul: `./kontrol.sh` (kitabı üretir, denetler, görüntüleyiciyi kurar). 8768 portu boş olmalı.
Liste SITE-HARITASI.md "Ekranlar" tablosuyla aynıdır.
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from ekran import adres, kaydet, sunucu, tarayici_ac, yeni_sayfa  # noqa: E402
from ortak import LAB, ekran_yolu  # noqa: E402
from terminal import calistir, ciz  # noqa: E402

KITAP = LAB / "defter" / "cikti" / "kahve-duragi-2026-01.xlsx"


def buyuk_defter_satiri(kod: int) -> int:
    """Büyük Defter sayfasında '<kod> · …' başlığının satırı (kitaptan okunur, elle yazılmaz)."""
    from openpyxl import load_workbook
    ws = load_workbook(KITAP)["Büyük Defter"]
    for r in range(1, ws.max_row + 1):
        v = ws[f"A{r}"].value
        if isinstance(v, str) and v.startswith(f"{kod} ·"):
            return r
    raise LookupError(kod)


# (ders, ad, sayfa, seçili hücre, vurgulanan aralık)
def tarayici_ekranlari() -> list[tuple[str, str, str, str | None, str | None]]:
    bd102 = buyuk_defter_satiri(102)
    bd100 = buyuk_defter_satiri(100)
    return [
        ("1.1", "bilanco-denklik", "bilanco", "B19", "A17:E17"),
        ("1.2", "yevmiye-ilk-maddeler", "yevmiye", "I5", "A5:I11"),
        ("1.3", "hesap-plani", "hesap-plani", "A5", None),
        ("1.3", "buyuk-defter-kasa", "buyuk-defter", f"F{bd100 + 3}", f"A{bd100}:G{bd100 + 4}"),
        ("1.4", "belge-sutunu", "yevmiye", "C5", "C5:C29"),
        ("2.1", "satis-belgeleri", "yevmiye", "C18", "A18:I23"),
        ("2.2", "kdv-ozet", "kdv", "D18", "A13:D19"),
        ("2.2", "kdv-ornek-ay", "kdv", "D27", "A21:D28"),
        ("2.3", "buyuk-defter-banka", "buyuk-defter", f"F{bd102 + 9}", f"A{bd102}:G{bd102 + 10}"),
        ("2.3", "banka-mutabakat", "banka-ekstresi", "E18", "A15:E19"),
        ("2.4", "mizan", "mizan", "E23", "A23:F23"),
        ("3.1", "stok-smm", "stok", "E15", "A12:E15"),
        ("3.2", "amortisman-tablosu", "amortisman", "B12", "A11:G17"),
        ("3.2", "dogrudan-gider-haddi", "amortisman", "C26", "A24:C27"),
        ("3.3", "gelir-tablosu", "gelir-tablosu", "B11", "A4:B11"),
        ("3.4", "bilanco", "bilanco", "E12", "A4:E17"),
        ("4.1", "vergi-takvimi", "takvim", "C5", "A5:D5"),
        ("4.4", "bilgi-sayfasi", "bilgi", "B15", "A14:C17"),
    ]


def terminal_ekranlari() -> list[tuple[str, str, list[str]]]:
    kontrol = ["uv", "run", "--project", "araclar", "python", "defter/kontrol.py"]
    return [
        ("1.2", "kontrol", kontrol),
        ("4.3", "negatif-test", kontrol + ["--negatif"]),
    ]


def terminal(ders: str, ad: str, komut: list[str]) -> None:
    cikti, _ = calistir(komut, LAB, None, None)
    cikti = cikti.replace(str(LAB), "~/lab")
    txt = ekran_yolu(ders, ad, ".txt")
    txt.write_text(f"~/lab $ {' '.join(komut)}\n{cikti}\n", encoding="utf-8")
    ciz(txt)


def main() -> None:
    secilen = set(sys.argv[1:])
    with sunucu() as taban, tarayici_ac() as t:
        s = yeni_sayfa(t)
        for ders, ad, sayfa, sec, vurgu in tarayici_ekranlari():
            if secilen and ders not in secilen:
                continue
            s.goto(adres(taban, sayfa, sec, vurgu), wait_until="networkidle")
            kaydet(s, ders, ad)
    for ders, ad, komut in terminal_ekranlari():
        if secilen and ders not in secilen:
            continue
        terminal(ders, ad, komut)


if __name__ == "__main__":
    main()
