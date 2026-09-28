"""Defter çalışma kitabını denetler: formüller hesaplanır (pycel), denklikler ve çıpa değerler kontrol edilir.

Kullanım (lab klasöründen):
    uv run --project araclar python defter/kontrol.py                         # Kahve Durağı kitabı + çıpalar
    uv run --project araclar python defter/kontrol.py dosya.xlsx              # başka bir kitap (çıpasız)
    uv run --project araclar python defter/kontrol.py --negatif               # bozuk kopyalarla negatif test

Kontroller:
  K1 her yevmiye maddesinde borç = alacak (hücre değerlerinden, formülden bağımsız)
  K2 yevmiye borç toplamı = alacak toplamı
  K3 mizan: hareket toplamları ve kalan toplamları denk; her kod hesap planında
  K4 bilanço: aktif toplamı = pasif toplamı
  K5 kitabın içindeki bütün kontrol hücreleri olumlu (DENK / EŞİT / TUTUYOR / AKTARILDI …)
  K6 hiçbir formül hata vermiyor (#VALUE!, #REF!, #NAME? …)
  K7 ÇIPA: sonuçlar, koddan bağımsız elle hesaplanmış beklenen.yml ile aynı (yalnız Kahve Durağı kitabı)

Denklik doğruluk demek değildir: borç ve alacak yer değiştirse de K1–K4 geçer; bunu K7 yakalar (--negatif).
"""

from __future__ import annotations

import shutil
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

import yaml
from openpyxl import load_workbook
from pycel import ExcelCompiler

DEFTER = Path(__file__).resolve().parent
KITAP = DEFTER / "cikti" / "kahve-duragi-2026-01.xlsx"
CAPA = DEFTER / "veri" / "beklenen.yml"
OLUMLU = {"DENK", "EŞİT", "TUTUYOR", "AKTARILDI", "doğrudan gider yazılabilir", "amortismana tabi"}
HATALAR = ("#VALUE!", "#REF!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")


def yuvarla(x) -> float:
    return round(float(x or 0), 2)


def tl(x: float) -> str:
    """1030400.0 → '1.030.400,00' (Türkçe sayı biçimi)."""
    return f"{x:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


class Denetim:
    def __init__(self, yol: Path, capa: bool):
        self.yol = yol
        self.capa = capa
        self.wb = load_workbook(yol)                    # formüller
        self.x = ExcelCompiler(filename=str(yol))      # hesaplanan değerler
        self.sonuc: list[tuple[str, bool, str]] = []

    def deger(self, adres: str):
        return self.x.evaluate(adres)

    def ekle(self, kod: str, ok: bool, ayrinti: str) -> None:
        self.sonuc.append((kod, ok, ayrinti))

    def yevmiye_satirlari(self):
        ws = self.wb["Yevmiye"]
        for r in range(5, ws.max_row + 1):
            m = ws[f"A{r}"].value
            if isinstance(m, (int, float)):
                yield r, int(m), ws[f"E{r}"].value, yuvarla(ws[f"G{r}"].value), yuvarla(ws[f"H{r}"].value)

    def k1_madde(self) -> None:
        toplam = defaultdict(lambda: [0.0, 0.0])
        for _, m, _, b, a in self.yevmiye_satirlari():
            toplam[m][0] += b
            toplam[m][1] += a
        bozuk = [f"madde {m}: borç {tl(b)} ≠ alacak {tl(a)}" for m, (b, a) in sorted(toplam.items()) if round(b - a, 2) != 0]
        self.ekle("K1 madde denkliği", not bozuk and bool(toplam),
                  (f"{len(toplam)} madde, hepsi denk" if toplam else "madde yok") if not bozuk else "; ".join(bozuk))

    def son_satir(self, ws, sutun: str, metin: str) -> int:
        for r in range(1, ws.max_row + 1):
            if ws[f"{sutun}{r}"].value == metin:
                return r
        raise LookupError(f"{ws.title}: '{metin}' satırı yok")

    def k2_yevmiye(self) -> None:
        t = self.son_satir(self.wb["Yevmiye"], "F", "TOPLAM")
        b, a = yuvarla(self.deger(f"Yevmiye!G{t}")), yuvarla(self.deger(f"Yevmiye!H{t}"))
        ham_b = round(sum(s[3] for s in self.yevmiye_satirlari()), 2)
        ok = b == a == ham_b
        self.ekle("K2 yevmiye toplamı", ok, f"borç {tl(b)} · alacak {tl(a)}" + ("" if b == ham_b else f" · hücre toplamı {tl(ham_b)}"))

    def k3_mizan(self) -> None:
        ws = self.wb["Mizan"]
        t = self.son_satir(ws, "B", "TOPLAM")
        c, d, e, f = (yuvarla(self.deger(f"Mizan!{s}{t}")) for s in "CDEF")
        kodlar = {ws[f"A{r}"].value for r in range(5, t)}
        eksik = sorted({k for _, _, k, _, _ in self.yevmiye_satirlari()} - kodlar)
        ok = c == d and e == f and not eksik
        self.ekle("K3 mizan dengesi", ok, f"hareket {tl(c)} = {tl(d)} · kalan {tl(e)} = {tl(f)}"
                  + (f" · hesap planında olmayan kod: {eksik}" if eksik else ""))

    def k4_bilanco(self) -> None:
        aktif, pasif = yuvarla(self.deger("Bilanço!B17")), yuvarla(self.deger("Bilanço!E17"))
        self.ekle("K4 bilanço denkliği", aktif == pasif, f"aktif {tl(aktif)} · pasif {tl(pasif)}")

    def k5_k6_hucreler(self) -> None:
        olumsuz, hatali, sayi = [], [], 0
        for ws in self.wb.worksheets:
            for satir in ws.iter_rows():
                for h in satir:
                    if not (isinstance(h.value, str) and h.value.startswith("=")):
                        continue
                    sayi += 1
                    v = self.deger(f"{ws.title}!{h.coordinate}")
                    if isinstance(v, str) and v in HATALAR:
                        hatali.append(f"{ws.title}!{h.coordinate} {v}")
                    elif '"DENK DEĞİL"' in h.value or '"FARK VAR"' in h.value or "EŞİT DEĞİL" in h.value or '"AKTARILMADI"' in h.value:
                        if v not in OLUMLU:
                            olumsuz.append(f"{ws.title}!{h.coordinate} = {v}")
        self.ekle("K5 kitap içi kontrol hücreleri", not olumsuz, "hepsi olumlu" if not olumsuz else "; ".join(olumsuz[:6]))
        self.ekle("K6 formül hataları", not hatali, f"{sayi} formül hesaplandı, hata yok" if not hatali else "; ".join(hatali[:6]))

    def k7_capa(self) -> None:
        beklenen = yaml.safe_load(CAPA.read_text(encoding="utf-8"))["degerler"]
        farklar = []
        for ad, v in beklenen.items():
            gercek = yuvarla(self.deger(v["hucre"]))
            if gercek != round(float(v["deger"]), 2):
                farklar.append(f"{ad}: beklenen {tl(float(v['deger']))} · kitap {tl(gercek)}")
        self.ekle("K7 çıpa (elle hesap)", not farklar, f"{len(beklenen)} değer tuttu" if not farklar else "; ".join(farklar))

    def calistir(self) -> bool:
        self.k1_madde()
        self.k2_yevmiye()
        self.k3_mizan()
        self.k4_bilanco()
        self.k5_k6_hucreler()
        if self.capa:
            self.k7_capa()
        return all(ok for _, ok, _ in self.sonuc)

    def yazdir(self) -> None:
        print(f"Dosya: {self.yol.name}")
        for kod, ok, ayrinti in self.sonuc:
            print(f"  {'GEÇTİ' if ok else 'KALDI':5}  {kod:30} {ayrinti}")


def negatif() -> bool:
    """Bilerek bozulmuş kopyalar: denetim her birini yakalamalı."""
    bozmalar = [
        ("Bir alacak tutarı yanlış yazıldı (madde 1: 300.000 → 30.000)",
         lambda ws: ws.__setitem__("H6", 30000), {"K1 madde denkliği", "K4 bilanço denkliği"}),
        ("Hesap planında olmayan kod (madde 2: 100 → 101)",
         lambda ws: ws.__setitem__("E7", 101), {"K3 mizan dengesi", "K5 kitap içi kontrol hücreleri"}),
        ("Borç ve alacak yer değiştirdi (madde 10: tahsilat ters kaydedildi) — denklik bozulmaz",
         lambda ws: (ws.__setitem__("G30", None), ws.__setitem__("H30", 48000),
                     ws.__setitem__("G31", 48000), ws.__setitem__("H31", None)), {"K7 çıpa (elle hesap)"}),
    ]
    tumu = True
    with tempfile.TemporaryDirectory() as gd:
        for ad, boz, beklenen in bozmalar:
            kopya = Path(gd) / "bozuk.xlsx"
            shutil.copy(KITAP, kopya)
            wb = load_workbook(kopya)
            boz(wb["Yevmiye"])
            wb.save(kopya)
            d = Denetim(kopya, capa=True)
            d.calistir()
            kalan = {k for k, ok, _ in d.sonuc if not ok}
            yakalandi = beklenen <= kalan
            tumu &= yakalandi
            print(f"  {'GEÇTİ' if yakalandi else 'KALDI':5}  {ad}")
            print(f"         KALDI veren kontroller: {', '.join(sorted(kalan)) or 'yok'}")
    return tumu


def main() -> int:
    if "--negatif" in sys.argv:
        print("Negatif test — bozuk kopyalar yakalanmalı:")
        return 0 if negatif() else 1
    arg = [a for a in sys.argv[1:] if not a.startswith("--")]
    yol = Path(arg[0]).resolve() if arg else KITAP
    d = Denetim(yol, capa=(yol == KITAP))
    ok = d.calistir()
    d.yazdir()
    print(f"SONUÇ: {'GEÇTİ' if ok else 'KALDI'}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
