"""Kahve Durağı (kurgusal) — Ocak 2026 defter çalışma kitabını üretir.

Kullanım (lab klasöründen):
    uv run --project araclar python defter/uret.py              # cikti/kahve-duragi-2026-01.xlsx
    uv run --project araclar python defter/uret.py --sablon     # cikti/bos-sablon.xlsx (yevmiye boş)

Girdi: defter/veri/{hesap-plani.csv, yevmiye.csv, stok-sayim.csv, ayarlar.yml}
Çıktı: gerçek .xlsx; bütün toplamlar, mizan, KDV özeti, amortisman, stok, gelir tablosu ve bilanço
FORMÜLDÜR (değer yazılmaz). Değerleri Excel / LibreOffice / Google E-Tablolar açarken hesaplar;
lab'da `kontrol.py` aynı formülleri pycel ile hesaplayıp denetler.

Bu dosya mali müşavirlik hizmeti değildir. Kendi durumunuz için mali müşavirinize danışın.
"""

from __future__ import annotations

import csv
import sys
from datetime import date
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DEFTER = Path(__file__).resolve().parent
VERI = DEFTER / "veri"
CIKTI = DEFTER / "cikti"

SON_SATIR = 200          # yevmiye formülleri bu satıra kadar okur (öğrenci satır ekleyebilsin)
PARA = '#,##0.00;(#,##0.00);"-"'
ORAN = '0%'
KALIN = Font(bold=True)
BASLIK = Font(bold=True, size=14)
NOT = Font(italic=True, size=9, color="555555")
BAS_DOLGU = PatternFill("solid", fgColor="E8E1D5")
TOP_DOLGU = PatternFill("solid", fgColor="F4F0E8")
KONTROL_DOLGU = PatternFill("solid", fgColor="E3F1E6")
INCE = Side(style="thin", color="BBBBBB")
UST = Border(top=Side(style="thin", color="333333"))
UYARI = "Bu çalışma kitabı eğitim amaçlıdır. Mali müşavirlik hizmeti değildir. Kendi durumunuz için mali müşavirinize danışın."


def oku_csv(ad: str) -> list[dict]:
    with open(VERI / ad, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def yaz(ws, hucre: str, deger, *, font=None, fmt=None, dolgu=None, hiza=None, kenar=None):
    c = ws[hucre]
    c.value = deger
    if font:
        c.font = font
    if fmt:
        c.number_format = fmt
    if dolgu:
        c.fill = dolgu
    if hiza:
        c.alignment = hiza
    if kenar:
        c.border = kenar
    return c


def basliklar(ws, satir: int, adlar: list[str]) -> None:
    for i, ad in enumerate(adlar, start=1):
        yaz(ws, f"{get_column_letter(i)}{satir}", ad, font=KALIN, dolgu=BAS_DOLGU,
            hiza=Alignment(horizontal="center", vertical="center", wrap_text=True))


def genislik(ws, olculer: dict[str, float]) -> None:
    for sutun, g in olculer.items():
        ws.column_dimensions[sutun].width = g


def sayfa_basi(ws, baslik: str, alt: str) -> None:
    yaz(ws, "A1", baslik, font=BASLIK)
    yaz(ws, "A2", alt, font=NOT)


# ---------------------------------------------------------------- sayfalar
def bilgi(wb, ayar: dict) -> None:
    ws = wb.active
    ws.title = "Bilgi"
    isl, don, par = ayar["isletme"], ayar["donem"], ayar["parametreler"]
    sayfa_basi(ws, "Kahve Durağı — Defter Çalışma Kitabı", "KURGUSAL İŞLETME · Muhasebe 101 lab'ı · " + UYARI)
    satirlar = [
        ("Unvan", isl["unvan"]), ("Faaliyet", isl["faaliyet"]), ("Vergi kimlik no", isl["vkn"]),
        ("Adres", isl["adres"]), ("Telefon", isl["telefon"]), ("E-posta", isl["eposta"]),
        ("Defter esası", isl["defter_esasi"]), ("Dönem", don["ad"]),
    ]
    for i, (k, v) in enumerate(satirlar, start=4):
        yaz(ws, f"A{i}", k, font=KALIN)
        yaz(ws, f"B{i}", v)
    yaz(ws, "A13", "Parametreler (formüller buradan okur)", font=KALIN)
    basliklar(ws, 14, ["Parametre", "Değer", "Kaynak / not"])
    # Sabit adresler: B15 KDV oranı, B16 doğrudan gider haddi, B17 örnek faydalı ömür
    for i, anahtar in enumerate(["kdv_genel_oran", "dogrudan_gider_haddi", "faydali_omur_ornek"], start=15):
        p = par[anahtar]
        yaz(ws, f"A{i}", p["etiket"])
        yaz(ws, f"B{i}", p["deger"], fmt=ORAN if anahtar == "kdv_genel_oran" else '#,##0')
        yaz(ws, f"C{i}", p["kaynak"] + (" · " + p["not"] if p.get("not") else ""),
            hiza=Alignment(wrap_text=True, vertical="top"))
        ws.row_dimensions[i].height = 48
    yaz(ws, "A19", "Sayfalar", font=KALIN)
    sayfalar = [
        ("Hesap Planı", "kullanılan hesaplar — Tekdüzen Hesap Planı kodlarıyla (MSUGT, RG 26.12.1992)"),
        ("Yevmiye", "tarih sırasıyla maddeler; her madde borç = alacak"),
        ("Büyük Defter", "yevmiyedeki satırların hesap hesap dağılımı ve bakiyesi"),
        ("Mizan", "hesapların borç/alacak toplamları ve kalanları; denklik kontrolü"),
        ("Banka Ekstresi", "bankanın belgesi (kurgusal) ve defterle mutabakat"),
        ("KDV", "hesaplanan, indirilecek, ödenecek / devreden KDV"),
        ("Stok", "ay sonu sayım → satılan malın maliyeti"),
        ("Amortisman", "demirbaşın yıllara dağıtılması; doğrudan gider haddi"),
        ("Gelir Tablosu", "dönemin kârı — mizandan formülle"),
        ("Bilanço", "ay sonu varlıklar = kaynaklar — mizandan formülle"),
        ("Takvim", "2026 vergi takvimi örnekleri (GİB, okundu 2026-09-25)"),
    ]
    for i, (s, a) in enumerate(sayfalar, start=20):
        yaz(ws, f"A{i}", s)
        yaz(ws, f"B{i}", a)
    genislik(ws, {"A": 34, "B": 58, "C": 90})


def hesap_plani(wb, hesaplar: list[dict]) -> None:
    ws = wb.create_sheet("Hesap Planı")
    sayfa_basi(ws, "Hesap Planı (kullanılan hesaplar)",
               "Kodlar Tekdüzen Hesap Planı'ndan (Muhasebe Sistemi Uygulama Genel Tebliği Sıra No: 1, RG 26.12.1992 sayı 21447 mük.). "
               "Ticaret işletmesi olarak giderler 7/B seçeneğiyle (79 Gider Çeşitleri) izlenir.")
    basliklar(ws, 4, ["Kod", "Hesap adı", "Grup", "Tür", "Normal bakiye"])
    for i, h in enumerate(hesaplar, start=5):
        yaz(ws, f"A{i}", int(h["kod"]))
        yaz(ws, f"B{i}", h["ad"])
        yaz(ws, f"C{i}", h["grup"])
        yaz(ws, f"D{i}", h["tur"])
        yaz(ws, f"E{i}", h["normal"])
    ws.freeze_panes = "A5"
    genislik(ws, {"A": 8, "B": 40, "C": 34, "D": 20, "E": 14})


def yevmiye(wb, satirlar: list[dict], son_hp: int) -> None:
    ws = wb.create_sheet("Yevmiye")
    sayfa_basi(ws, "Yevmiye Defteri — Ocak 2026",
               "Her madde bir belgeye dayanır; bir maddenin borç toplamı alacak toplamına eşit olmalıdır (I sütunu).")
    basliklar(ws, 4, ["Madde", "Tarih", "Belge", "Açıklama", "Hesap", "Hesap adı", "Borç", "Alacak", "Madde kontrolü"])
    r = f"$5:$" + str(SON_SATIR)
    onceki = None
    for i, s in enumerate(satirlar, start=5):
        m = int(s["madde"])
        yaz(ws, f"A{i}", m)
        yaz(ws, f"B{i}", date.fromisoformat(s["tarih"]), fmt="DD.MM.YYYY")
        yaz(ws, f"C{i}", s["belge"])
        yaz(ws, f"D{i}", s["aciklama"])
        yaz(ws, f"E{i}", int(s["hesap"]))
        yaz(ws, f"F{i}", f"=IFERROR(VLOOKUP(E{i},'Hesap Planı'!$A$5:$B${son_hp},2,FALSE),\"HESAP PLANINDA YOK\")")
        if s["borc"]:
            yaz(ws, f"G{i}", float(s["borc"]), fmt=PARA)
        if s["alacak"]:
            yaz(ws, f"H{i}", float(s["alacak"]), fmt=PARA)
        if m != onceki:   # maddenin ilk satırında kontrol
            yaz(ws, f"I{i}",
                f"=IF(ROUND(SUMIF($A$5:$A${SON_SATIR},A{i},$G$5:$G${SON_SATIR})-SUMIF($A$5:$A${SON_SATIR},A{i},$H$5:$H${SON_SATIR}),2)=0,\"DENK\",\"DENK DEĞİL\")",
                dolgu=KONTROL_DOLGU)
            if onceki is not None:
                for sut in "ABCDEFGHI":
                    ws[f"{sut}{i}"].border = Border(top=INCE)
        onceki = m
    for i in range(5 + len(satirlar), SON_SATIR + 1):   # boş satırlarda hesap adı formülü hazır dursun
        yaz(ws, f"F{i}", f"=IF(E{i}=\"\",\"\",IFERROR(VLOOKUP(E{i},'Hesap Planı'!$A$5:$B${son_hp},2,FALSE),\"HESAP PLANINDA YOK\"))")
    t = SON_SATIR + 1
    yaz(ws, f"F{t}", "TOPLAM", font=KALIN, dolgu=TOP_DOLGU)
    yaz(ws, f"G{t}", f"=SUM(G5:G{SON_SATIR})", font=KALIN, fmt=PARA, dolgu=TOP_DOLGU, kenar=UST)
    yaz(ws, f"H{t}", f"=SUM(H5:H{SON_SATIR})", font=KALIN, fmt=PARA, dolgu=TOP_DOLGU, kenar=UST)
    yaz(ws, f"I{t}", f"=IF(ROUND(G{t}-H{t},2)=0,\"DENK\",\"DENK DEĞİL\")", font=KALIN, dolgu=KONTROL_DOLGU)
    ws.freeze_panes = "A5"
    genislik(ws, {"A": 8, "B": 12, "C": 26, "D": 44, "E": 8, "F": 30, "G": 14, "H": 14, "I": 14})


def buyuk_defter(wb, satirlar: list[dict], hesaplar: list[dict]) -> None:
    ws = wb.create_sheet("Büyük Defter")
    sayfa_basi(ws, "Büyük Defter (Defteri Kebir) — Ocak 2026",
               "Yevmiyedeki her satır kendi hesabına taşınır; bakiye = borç − alacak (artı: borç bakiyesi, eksi: alacak bakiyesi).")
    ad = {h["kod"]: h["ad"] for h in hesaplar}
    r = 4
    for h in hesaplar:
        kod = h["kod"]
        idx = [i for i, s in enumerate(satirlar, start=5) if s["hesap"] == kod]
        if not idx:
            continue
        yaz(ws, f"A{r}", f"{kod} · {ad[kod]}", font=Font(bold=True, size=12))
        r += 1
        basliklar(ws, r, ["Tarih", "Madde", "Açıklama", "Borç", "Alacak", "Bakiye", "Yön"])
        r += 1
        ilk = r
        for j, y in enumerate(idx):
            yaz(ws, f"A{r}", f"=Yevmiye!B{y}", fmt="DD.MM.YYYY")
            yaz(ws, f"B{r}", f"=Yevmiye!A{y}")
            yaz(ws, f"C{r}", f"=Yevmiye!D{y}")
            yaz(ws, f"D{r}", f"=N(Yevmiye!G{y})", fmt=PARA)
            yaz(ws, f"E{r}", f"=N(Yevmiye!H{y})", fmt=PARA)
            onceki = "0" if j == 0 else f"F{r-1}"
            yaz(ws, f"F{r}", f"={onceki}+D{r}-E{r}", fmt=PARA)
            yaz(ws, f"G{r}", f"=IF(ROUND(F{r},2)>0,\"Borç\",IF(ROUND(F{r},2)<0,\"Alacak\",\"—\"))")
            r += 1
        yaz(ws, f"C{r}", "Toplam", font=KALIN, dolgu=TOP_DOLGU)
        yaz(ws, f"D{r}", f"=SUM(D{ilk}:D{r-1})", font=KALIN, fmt=PARA, dolgu=TOP_DOLGU, kenar=UST)
        yaz(ws, f"E{r}", f"=SUM(E{ilk}:E{r-1})", font=KALIN, fmt=PARA, dolgu=TOP_DOLGU, kenar=UST)
        yaz(ws, f"F{r}", f"=D{r}-E{r}", font=KALIN, fmt=PARA, dolgu=TOP_DOLGU, kenar=UST)
        r += 2
    genislik(ws, {"A": 12, "B": 8, "C": 56, "D": 14, "E": 14, "F": 14, "G": 9})


def mizan(wb, hesaplar: list[dict]) -> int:
    ws = wb.create_sheet("Mizan")
    sayfa_basi(ws, "Mizan — 31.01.2026",
               "Toplamlar Yevmiye sayfasından SUMIF ile gelir. Borç toplamı = alacak toplamı ve borç kalanı = alacak kalanı olmalıdır.")
    basliklar(ws, 4, ["Kod", "Hesap adı", "Borç toplamı", "Alacak toplamı", "Borç kalanı", "Alacak kalanı"])
    Y = f"Yevmiye!$E$5:$E${SON_SATIR}"
    for i, h in enumerate(hesaplar, start=5):
        yaz(ws, f"A{i}", int(h["kod"]))
        yaz(ws, f"B{i}", h["ad"])
        yaz(ws, f"C{i}", f"=SUMIF({Y},A{i},Yevmiye!$G$5:$G${SON_SATIR})", fmt=PARA)
        yaz(ws, f"D{i}", f"=SUMIF({Y},A{i},Yevmiye!$H$5:$H${SON_SATIR})", fmt=PARA)
        yaz(ws, f"E{i}", f"=MAX(C{i}-D{i},0)", fmt=PARA)
        yaz(ws, f"F{i}", f"=MAX(D{i}-C{i},0)", fmt=PARA)
    son = 4 + len(hesaplar)
    t = son + 1
    yaz(ws, f"B{t}", "TOPLAM", font=KALIN, dolgu=TOP_DOLGU)
    for s in "CDEF":
        yaz(ws, f"{s}{t}", f"=SUM({s}5:{s}{son})", font=KALIN, fmt=PARA, dolgu=TOP_DOLGU, kenar=UST)
    k = t + 2
    yaz(ws, f"A{k}", "Kontroller", font=KALIN)
    kontroller = [
        ("Hareket toplamları (borç = alacak)", f"=IF(ROUND(C{t}-D{t},2)=0,\"DENK\",\"DENK DEĞİL\")"),
        ("Kalan toplamları (borç kalanı = alacak kalanı)", f"=IF(ROUND(E{t}-F{t},2)=0,\"DENK\",\"DENK DEĞİL\")"),
        ("Mizan borç toplamı = Yevmiye borç toplamı",
         f"=IF(ROUND(C{t}-Yevmiye!G{SON_SATIR+1},2)=0,\"EŞİT\",\"EŞİT DEĞİL — hesap planında olmayan kod var mı?\")"),
        ("7/B gider çeşitleri aktarıldı mı (79x − 798 = 0)",
         f"=IF(ROUND(SUMIF(A5:A{son},\">=790\",E5:E{son})-SUMIF(A5:A{son},\">=790\",F5:F{son}),2)=0,\"AKTARILDI\",\"AKTARILMADI\")"),
    ]
    for i, (ad, f) in enumerate(kontroller, start=k + 1):
        yaz(ws, f"A{i}", ad)
        yaz(ws, f"E{i}", f, font=KALIN, dolgu=KONTROL_DOLGU)
    ws.freeze_panes = "A5"
    genislik(ws, {"A": 8, "B": 40, "C": 15, "D": 15, "E": 15, "F": 15})
    return son


def kdv(wb, satirlar: list[dict]) -> None:
    ws = wb.create_sheet("KDV")
    sayfa_basi(ws, "KDV — Ocak 2026",
               "Oran Bilgi!B15'ten okunur (genel oran, kaynaklı). Hesaplanan KDV satışta 391'e, indirilecek KDV alışta 191'e yazılır.")
    basliklar(ws, 4, ["Madde", "Belge", "Tür", "Matrah", "Oran", "KDV (hesap)", "Deftere yazılan", "Kontrol"])
    belgeler: dict[int, dict] = {}
    for s in satirlar:
        m = int(s["madde"])
        if s["hesap"] in ("391", "191") and m != 15:
            b = belgeler.setdefault(m, {"belge": s["belge"], "tur": "Satış" if s["hesap"] == "391" else "Alış"})
    for s in satirlar:   # matrah: KDV satırı dışındaki gelir/gider/varlık satırı
        m = int(s["madde"])
        if m in belgeler and s["hesap"] in ("600", "255", "794", "153", "793"):
            belgeler[m]["matrah"] = float(s["borc"] or s["alacak"])
    r = 5
    for m, b in belgeler.items():
        hesap = 391 if b["tur"] == "Satış" else 191
        sut = "H" if b["tur"] == "Satış" else "G"
        yaz(ws, f"A{r}", m)
        yaz(ws, f"B{r}", b["belge"])
        yaz(ws, f"C{r}", b["tur"])
        yaz(ws, f"D{r}", b["matrah"], fmt=PARA)
        yaz(ws, f"E{r}", "=Bilgi!$B$15", fmt=ORAN)
        yaz(ws, f"F{r}", f"=ROUND(D{r}*E{r},2)", fmt=PARA)
        yaz(ws, f"G{r}", f"=SUMIFS(Yevmiye!${sut}$5:${sut}${SON_SATIR},Yevmiye!$A$5:$A${SON_SATIR},A{r},Yevmiye!$E$5:$E${SON_SATIR},{hesap})", fmt=PARA)
        yaz(ws, f"H{r}", f"=IF(ROUND(F{r}-G{r},2)=0,\"TUTUYOR\",\"FARK VAR\")", dolgu=KONTROL_DOLGU)
        r += 1
    son = r - 1
    r += 1
    yaz(ws, f"A{r}", "Dönem özeti", font=KALIN)
    ozet = r + 1
    satir = [
        ("Hesaplanan KDV (satışlar)", f"=SUMIF(C5:C{son},\"Satış\",F5:F{son})"),
        ("İndirilecek KDV (alışlar)", f"=SUMIF(C5:C{son},\"Alış\",F5:F{son})"),
        ("Önceki dönemden devreden KDV", 0),
        ("Ödenecek KDV", f"=MAX(D{ozet}-D{ozet+1}-D{ozet+2},0)"),
        ("Sonraki döneme devreden KDV", f"=MAX(D{ozet+1}+D{ozet+2}-D{ozet},0)"),
        ("Kontrol: devreden = 191 hesabının ay sonu borç kalanı",
         f"=IF(ROUND(D{ozet+4}-SUMIF(Mizan!A5:A40,191,Mizan!E5:E40),2)=0,\"TUTUYOR\",\"FARK VAR\")"),
    ]
    for i, (ad, f) in enumerate(satir, start=ozet):
        yaz(ws, f"B{i}", ad, font=KALIN if i in (ozet + 3, ozet + 4) else None)
        yaz(ws, f"D{i}", f, fmt=PARA if i < ozet + 5 else None,
            dolgu=KONTROL_DOLGU if i == ozet + 5 else None)
    r = ozet + 7
    yaz(ws, f"A{r}", "ÖRNEK ay — rakamlar uydurmadır (devreden KDV'nin nasıl kullanıldığını görmek için)", font=KALIN)
    o = r + 1
    ornek = [
        ("Satış matrahı (ÖRNEK)", 300000),
        ("Alış matrahı (ÖRNEK)", 120000),
        ("Hesaplanan KDV", f"=ROUND(D{o}*Bilgi!$B$15,2)"),
        ("İndirilecek KDV", f"=ROUND(D{o+1}*Bilgi!$B$15,2)"),
        ("Önceki dönemden devreden (Ocak'tan)", f"=D{ozet+4}"),
        ("Ödenecek KDV", f"=MAX(D{o+2}-D{o+3}-D{o+4},0)"),
        ("Sonraki döneme devreden", f"=MAX(D{o+3}+D{o+4}-D{o+2},0)"),
    ]
    for i, (ad, f) in enumerate(ornek, start=o):
        yaz(ws, f"B{i}", ad)
        yaz(ws, f"D{i}", f, fmt=PARA)
    r = o + len(ornek) + 1
    yaz(ws, f"A{r}", "KDV dahil tutardan matrahı bulmak", font=KALIN)
    d = r + 1
    yaz(ws, f"B{d}", "KDV dahil tutar (ÖRNEK)")
    yaz(ws, f"D{d}", 1200, fmt=PARA)
    yaz(ws, f"B{d+1}", "Matrah = dahil tutar / (1 + oran)")
    yaz(ws, f"D{d+1}", f"=ROUND(D{d}/(1+Bilgi!$B$15),2)", fmt=PARA)
    yaz(ws, f"B{d+2}", "KDV = dahil tutar − matrah")
    yaz(ws, f"D{d+2}", f"=D{d}-D{d+1}", fmt=PARA)
    yaz(ws, f"A{d+4}", "Beyan ve ödeme tarihleri kanunda değil, GİB vergi takviminde güncel hâliyle görülür (ör. Ağustos 2026 dönemi: son gün 28.09.2026).", font=NOT)
    genislik(ws, {"A": 8, "B": 44, "C": 8, "D": 15, "E": 7, "F": 14, "G": 15, "H": 12})


def stok(wb, sayim: list[dict]) -> None:
    ws = wb.create_sheet("Stok")
    sayfa_basi(ws, "Stok ve Satılan Malın Maliyeti — Ocak 2026",
               "Aralıklı envanter: ay sonunda sayılır; SMM = dönem başı stok + dönem içi alışlar − dönem sonu stok.")
    basliklar(ws, 4, ["Ürün (ay sonu sayım)", "Birim", "Miktar", "Birim maliyet", "Tutar"])
    for i, s in enumerate(sayim, start=5):
        yaz(ws, f"A{i}", s["urun"])
        yaz(ws, f"B{i}", s["birim"])
        yaz(ws, f"C{i}", int(s["miktar"]))
        yaz(ws, f"D{i}", float(s["birim_maliyet"]), fmt=PARA)
        yaz(ws, f"E{i}", f"=C{i}*D{i}", fmt=PARA)
    son = 4 + len(sayim)
    t = son + 1
    yaz(ws, f"A{t}", "Dönem sonu stok (sayım toplamı)", font=KALIN, dolgu=TOP_DOLGU)
    yaz(ws, f"E{t}", f"=SUM(E5:E{son})", font=KALIN, fmt=PARA, dolgu=TOP_DOLGU, kenar=UST)
    h = t + 2
    yaz(ws, f"A{h}", "SMM hesabı", font=KALIN)
    Y = f"Yevmiye!$E$5:$E${SON_SATIR}"
    hesap = [
        ("Dönem başı stok", 0),
        ("+ Dönem içi alışlar (153 borç, sayım kaydı hariç)",
         f"=SUMIFS(Yevmiye!$G$5:$G${SON_SATIR},{Y},153,Yevmiye!$A$5:$A${SON_SATIR},\"<>13\")"),
        ("− Dönem sonu stok (sayım)", f"=E{t}"),
        ("Satılan malın maliyeti (SMM)", f"=E{h+1}+E{h+2}-E{h+3}"),
        ("Kontrol: SMM = yevmiyedeki 621 kaydı", f"=IF(ROUND(E{h+4}-SUMIF({Y},621,Yevmiye!$G$5:$G${SON_SATIR}),2)=0,\"TUTUYOR\",\"FARK VAR\")"),
        ("Kontrol: sayım = 153 hesabının ay sonu kalanı", f"=IF(ROUND(E{t}-SUMIF(Mizan!A5:A40,153,Mizan!E5:E40),2)=0,\"TUTUYOR\",\"FARK VAR\")"),
    ]
    for i, (ad, f) in enumerate(hesap, start=h + 1):
        yaz(ws, f"A{i}", ad, font=KALIN if "SMM)" in ad else None)
        yaz(ws, f"E{i}", f, fmt=PARA if "Kontrol" not in ad else None,
            dolgu=KONTROL_DOLGU if "Kontrol" in ad else None)
    genislik(ws, {"A": 50, "B": 8, "C": 9, "D": 14, "E": 15})


def amortisman(wb) -> None:
    ws = wb.create_sheet("Amortisman")
    sayfa_basi(ws, "Amortisman — Tezgâh ve raf sistemi",
               "VUK md. 313–320: normal amortisman oranı = 1 / faydalı ömür; azalan bakiyeler oranı normalin iki katı, en çok %50 (VUK mük. md. 315).")
    Y = f"Yevmiye!$E$5:$E${SON_SATIR}"
    girdi = [
        ("Kıymet", "Tezgâh ve raf sistemi (demirbaş, 255)"),
        ("Maliyet bedeli (255 borç)", f"=SUMIF({Y},255,Yevmiye!$G$5:$G${SON_SATIR})"),
        ("Aktife giriş", date(2026, 1, 6)),
        ("Faydalı ömür (yıl) — ÖRNEK", "=Bilgi!$B$17"),
        ("Normal amortisman oranı", "=1/B7"),
        ("Azalan bakiyeler oranı", "=MIN(2*B8,0.5)"),
    ]
    for i, (ad, f) in enumerate(girdi, start=4):
        yaz(ws, f"A{i}", ad, font=KALIN)
        c = yaz(ws, f"B{i}", f)
        if i == 5:
            c.number_format = PARA
        elif i == 6:
            c.number_format = "DD.MM.YYYY"
        elif i in (8, 9):
            c.number_format = ORAN
    basliklar(ws, 11, ["Yıl", "Normal: amortisman", "Normal: birikmiş", "Normal: net değer",
                       "Azalan: dönem başı", "Azalan: amortisman", "Azalan: net değer"])
    for j in range(5):
        r = 12 + j
        yaz(ws, f"A{r}", 2026 + j)
        yaz(ws, f"B{r}", f"=IF(A{r}-$A$12<$B$7,ROUND($B$5*$B$8,2),0)", fmt=PARA)
        yaz(ws, f"C{r}", f"=SUM($B$12:B{r})", fmt=PARA)
        yaz(ws, f"D{r}", f"=$B$5-C{r}", fmt=PARA)
        yaz(ws, f"E{r}", "=$B$5" if j == 0 else f"=G{r-1}", fmt=PARA)
        yaz(ws, f"F{r}", f"=IF(A{r}-$A$12=$B$7-1,E{r},ROUND(E{r}*$B$9,2))", fmt=PARA)
        yaz(ws, f"G{r}", f"=E{r}-F{r}", fmt=PARA)
    yaz(ws, "A17", "Toplam", font=KALIN, dolgu=TOP_DOLGU)
    yaz(ws, "B17", "=SUM(B12:B16)", font=KALIN, fmt=PARA, dolgu=TOP_DOLGU, kenar=UST)
    yaz(ws, "F17", "=SUM(F12:F16)", font=KALIN, fmt=PARA, dolgu=TOP_DOLGU, kenar=UST)
    yaz(ws, "A18", "Kontrol: her iki yöntemde toplam = maliyet", font=KALIN)
    yaz(ws, "B18", "=IF(AND(ROUND(B17-B5,2)=0,ROUND(F17-B5,2)=0),\"TUTUYOR\",\"FARK VAR\")", dolgu=KONTROL_DOLGU)
    yaz(ws, "A20", "Ocak payı (ara dönem, yıllık / 12)", font=KALIN)
    yaz(ws, "B20", "=ROUND(B12/12,2)", fmt=PARA)
    yaz(ws, "A21", "Kontrol: yevmiyedeki 796 kaydı", font=KALIN)
    yaz(ws, "B21", f"=IF(ROUND(B20-SUMIF({Y},796,Yevmiye!$G$5:$G${SON_SATIR}),2)=0,\"TUTUYOR\",\"FARK VAR\")", dolgu=KONTROL_DOLGU)
    yaz(ws, "A22", "Not: vergi hesabında amortisman yıllık ayrılır (VUK md. 320); aylık pay yalnız ara dönem yönetim tablosu içindir.", font=NOT)
    yaz(ws, "A24", "Doğrudan gider mi, amortisman mı? (VUK md. 313 — 2026 haddi Bilgi!B16)", font=KALIN)
    basliklar(ws, 25, ["Kıymet", "Değer", "Sonuç"])
    for r, (ad, deger) in enumerate([("Barkod okuyucu", 8000), ("Tezgâh ve raf sistemi", "=B5")], start=26):
        yaz(ws, f"A{r}", ad)
        yaz(ws, f"B{r}", deger, fmt=PARA)
        yaz(ws, f"C{r}", f"=IF(B{r}<=Bilgi!$B$16,\"doğrudan gider yazılabilir\",\"amortismana tabi\")", dolgu=KONTROL_DOLGU)
    genislik(ws, {"A": 40, "B": 30, "C": 26, "D": 16, "E": 18, "F": 18, "G": 16})


def gelir_tablosu(wb) -> None:
    ws = wb.create_sheet("Gelir Tablosu")
    sayfa_basi(ws, "Gelir Tablosu — 01.01.2026–31.01.2026 (ara dönem)",
               "Tutarlar Mizan sayfasından formülle gelir. Vergi karşılığı bu çalışmanın kapsamı dışında (vergi öncesi kâr).")
    M = "Mizan!$A$5:$A$40"
    satir = [
        ("A. Brüt satışlar — Yurt içi satışlar (600)", f"=SUMIF({M},600,Mizan!$F$5:$F$40)-SUMIF({M},600,Mizan!$E$5:$E$40)", False),
        ("B. Satış indirimleri (−)", 0, False),
        ("C. Net satışlar", "=B4-B5", True),
        ("D. Satışların maliyeti (−) — Satılan ticari mallar maliyeti (621)", f"=SUMIF({M},621,Mizan!$E$5:$E$40)-SUMIF({M},621,Mizan!$F$5:$F$40)", False),
        ("Brüt satış kârı", "=B6-B7", True),
        ("E. Faaliyet giderleri (−) — Genel yönetim giderleri (632)", f"=SUMIF({M},632,Mizan!$E$5:$E$40)-SUMIF({M},632,Mizan!$F$5:$F$40)", False),
        ("Faaliyet kârı", "=B8-B9", True),
        ("Dönem kârı (vergi öncesi)", "=B10", True),
    ]
    for i, (ad, f, kalin) in enumerate(satir, start=4):
        yaz(ws, f"A{i}", ad, font=KALIN if kalin else None, dolgu=TOP_DOLGU if kalin else None)
        yaz(ws, f"B{i}", f, fmt=PARA, font=KALIN if kalin else None, dolgu=TOP_DOLGU if kalin else None)
    yaz(ws, "A13", "Brüt kâr marjı", font=KALIN)
    yaz(ws, "B13", "=IF(B6=0,0,B8/B6)", fmt="0.0%")
    yaz(ws, "A15", "Kontrol: gelir tablosu hesaplarının net kalanı = dönem kârı", font=KALIN)
    yaz(ws, "B15", f"=IF(ROUND(SUMIFS(Mizan!$F$5:$F$40,{M},\">=600\",{M},\"<700\")-SUMIFS(Mizan!$E$5:$E$40,{M},\">=600\",{M},\"<700\")-B11,2)=0,\"TUTUYOR\",\"FARK VAR\")",
        dolgu=KONTROL_DOLGU)
    genislik(ws, {"A": 64, "B": 18})


def bilanco(wb) -> None:
    ws = wb.create_sheet("Bilanço")
    sayfa_basi(ws, "Bilanço — 31.01.2026 (ara dönem)",
               "Aktif (varlıklar) = Pasif (kaynaklar). Dönem kârı Gelir Tablosu'ndan gelir; resmî yıl sonu bilançosu değildir.")
    M = "Mizan!$A$5:$A$40"

    # net kalan: aktif kalemi = borç kalanı − alacak kalanı; pasif kalemi = alacak kalanı − borç kalanı.
    # Böylece ters bakiye veren bir hata (ör. 391 borç bakiyesi) bilançoya eksi olarak yansır, denklik bozulur.
    def b(kod: int) -> str:
        return f"(SUMIF({M},{kod},Mizan!$E$5:$E$40)-SUMIF({M},{kod},Mizan!$F$5:$F$40))"

    def a(kod: int) -> str:
        return f"(SUMIF({M},{kod},Mizan!$F$5:$F$40)-SUMIF({M},{kod},Mizan!$E$5:$E$40))"

    basliklar(ws, 4, ["AKTİF (Varlıklar)", "Tutar", "", "PASİF (Kaynaklar)", "Tutar"])
    aktif = [
        ("I. DÖNEN VARLIKLAR", None, True),
        ("  Kasa (100)", f"={b(100)}", False),
        ("  Bankalar (102)", f"={b(102)}", False),
        ("  Alıcılar (120)", f"={b(120)}", False),
        ("  Ticari mallar (153)", f"={b(153)}", False),
        ("  İndirilecek KDV — devreden (191)", f"={b(191)}", False),
        ("  Dönen varlıklar toplamı", "=SUM(B6:B10)", True),
        ("II. DURAN VARLIKLAR", None, True),
        ("  Demirbaşlar (255)", f"={b(255)}", False),
        ("  Birikmiş amortismanlar (257) (−)", f"=-{a(257)}", False),
        ("  Duran varlıklar toplamı", "=SUM(B13:B14)", True),
    ]
    pasif = [
        ("III. KISA VADELİ YABANCI KAYNAKLAR", None, True),
        ("  Satıcılar (320)", f"={a(320)}", False),
        ("  Ödenecek vergi ve fonlar (360)", f"={a(360)}", False),
        ("  Hesaplanan KDV (391)", f"={a(391)}", False),
        ("  Kısa vadeli yabancı kaynaklar toplamı", "=SUM(E6:E8)", True),
        ("V. ÖZ KAYNAKLAR", None, True),
        ("  Sermaye (500)", f"={a(500)}", False),
        ("  Dönem kârı (Gelir Tablosu, vergi öncesi)", "='Gelir Tablosu'!B11", False),
        ("  Öz kaynaklar toplamı", "=SUM(E11:E12)", True),
    ]
    for i, (ad, f, kalin) in enumerate(aktif, start=5):
        yaz(ws, f"A{i}", ad, font=KALIN if kalin else None)
        if f:
            yaz(ws, f"B{i}", f, fmt=PARA, font=KALIN if kalin else None, dolgu=TOP_DOLGU if kalin else None)
    for i, (ad, f, kalin) in enumerate(pasif, start=5):
        yaz(ws, f"D{i}", ad, font=KALIN if kalin else None)
        if f:
            yaz(ws, f"E{i}", f, fmt=PARA, font=KALIN if kalin else None, dolgu=TOP_DOLGU if kalin else None)
    yaz(ws, "A17", "AKTİF TOPLAMI", font=KALIN, dolgu=BAS_DOLGU)
    yaz(ws, "B17", "=B11+B15", font=KALIN, fmt=PARA, dolgu=BAS_DOLGU, kenar=UST)
    yaz(ws, "D17", "PASİF TOPLAMI", font=KALIN, dolgu=BAS_DOLGU)
    yaz(ws, "E17", "=E9+E13", font=KALIN, fmt=PARA, dolgu=BAS_DOLGU, kenar=UST)
    yaz(ws, "A19", "Kontrol: aktif = pasif", font=KALIN)
    yaz(ws, "B19", "=IF(ROUND(B17-E17,2)=0,\"DENK\",\"DENK DEĞİL\")", font=KALIN, dolgu=KONTROL_DOLGU)
    genislik(ws, {"A": 40, "B": 16, "C": 3, "D": 44, "E": 16})


def banka_ekstresi(wb, satirlar: list[dict]) -> None:
    ws = wb.create_sheet("Banka Ekstresi")
    sayfa_basi(ws, "Banka Ekstresi — Örnek Bank (KURGUSAL) · Ocak 2026",
               "Bankanın gönderdiği belge (dış belge): tutarlar elle yazılıdır, defterden formülle gelmez. Mutabakat: ekstre bakiyesi = defterdeki 102 bakiyesi mi?")
    basliklar(ws, 4, ["Tarih", "Açıklama (bankanın yazdığı)", "Giren", "Çıkan", "Bakiye"])
    hareket = [(s["tarih"], s["aciklama"], float(s["borc"] or 0), float(s["alacak"] or 0))
               for s in satirlar if s["hesap"] == "102"]
    hareket.append(("2026-01-31", "Hesap işletim ücreti", 0.0, 50.0))   # deftere işlenmemiş banka masrafı
    r = 5
    for j, (t, a, g, c) in enumerate(hareket):
        yaz(ws, f"A{r}", date.fromisoformat(t), fmt="DD.MM.YYYY")
        yaz(ws, f"B{r}", a)
        if g:
            yaz(ws, f"C{r}", g, fmt=PARA)
        if c:
            yaz(ws, f"D{r}", c, fmt=PARA)
        yaz(ws, f"E{r}", ("=0" if j == 0 else f"=E{r-1}") + f"+N(C{r})-N(D{r})", fmt=PARA)
        r += 1
    son = r - 1
    m = r + 1
    yaz(ws, f"A{m}", "Mutabakat", font=KALIN)
    yaz(ws, f"B{m+1}", "Ekstre bakiyesi (31.01.2026)")
    yaz(ws, f"E{m+1}", f"=E{son}", fmt=PARA)
    yaz(ws, f"B{m+2}", "Defterdeki 102 Bankalar bakiyesi (Mizan)")
    yaz(ws, f"E{m+2}", "=SUMIF(Mizan!$A$5:$A$40,102,Mizan!$E$5:$E$40)", fmt=PARA)
    yaz(ws, f"B{m+3}", "Fark (ekstre − defter)", font=KALIN)
    yaz(ws, f"E{m+3}", f"=E{m+1}-E{m+2}", fmt=PARA, font=KALIN, dolgu=TOP_DOLGU)
    yaz(ws, f"B{m+4}", "Farkın nedeni: bankanın kestiği ücret deftere henüz işlenmedi (Ders 2.3 görevi). "
        "Hangi hesaba yazılacağını mali müşavirinle netleştir.", font=NOT)
    genislik(ws, {"A": 12, "B": 54, "C": 14, "D": 14, "E": 15})


def takvim(wb) -> None:
    ws = wb.create_sheet("Takvim")
    sayfa_basi(ws, "Vergi Takvimi Örnekleri — 2026",
               "Kaynak: GİB 2026 yıllık vergi takvimi ve gib.gov.tr vergi takvimi sayfası (okundu 2026-09-25). Kendi yükümlülüklerin için güncel takvime ve mali müşavirine bak.")
    basliklar(ws, 4, ["Dönem", "Yükümlülük", "Son gün", "Not"])
    satirlar = [
        ("Ocak 2026", "KDV beyanı ve ödemesi", date(2026, 3, 2), "Kahve Durağı'nın Ocak beyanı; 28 Şubat Cumartesi → ilk iş günü"),
        ("Ocak 2026", "Muhtasar ve prim hizmet beyannamesi", date(2026, 2, 26), "çalışanı ya da stopaj kesintisi olan işletmeler"),
        ("Şubat 2026", "KDV beyanı ve ödemesi", date(2026, 3, 30), "28 Mart Cumartesi → 30 Mart Pazartesi"),
        ("2025 yılı", "Yıllık gelir vergisi beyanı ve 1. taksit", date(2026, 3, 31), "şahıs işletmeleri (gelir vergisi mükellefleri)"),
        ("2025 yılı", "Kurumlar vergisi beyanı ve ödemesi", date(2026, 4, 30), "şirketler (Kahve Durağı Ltd. Şti. gibi)"),
        ("2026 I. dönem (Ocak–Mart)", "Geçici vergi beyanı ve ödemesi", date(2026, 5, 18), "17 Mayıs Pazar → 18 Mayıs"),
        ("Ağustos 2026", "KDV beyanı ve ödemesi", date(2026, 9, 28), "gib.gov.tr vergi takvimi sayfası, 2026-09-25"),
        ("Ağustos 2026", "Muhtasar ve prim hizmet beyannamesi", date(2026, 9, 28), "26 Eylül Cumartesi → 28 Eylül Pazartesi"),
    ]
    for i, (d, y, g, n) in enumerate(satirlar, start=5):
        yaz(ws, f"A{i}", d)
        yaz(ws, f"B{i}", y)
        yaz(ws, f"C{i}", g, fmt="DD.MM.YYYY")
        yaz(ws, f"D{i}", n)
    yaz(ws, "A14", "Kanun metnindeki gün ile takvimdeki gün aynı olmayabilir: KDV Kanunu md. 41 ve 46 beyan için 24., ödeme için 26. günü yazar;", font=NOT)
    yaz(ws, "A15", "2026 takviminde KDV son günü 28'idir (hafta sonuna denk gelirse ilk iş günü). Her zaman güncel GİB takvimine bak.", font=NOT)
    genislik(ws, {"A": 26, "B": 42, "C": 12, "D": 60})


def uret(sablon: bool = False) -> Path:
    ayar = yaml.safe_load((VERI / "ayarlar.yml").read_text(encoding="utf-8"))
    hesaplar = oku_csv("hesap-plani.csv")
    satirlar = [] if sablon else oku_csv("yevmiye.csv")
    sayim = oku_csv("stok-sayim.csv")
    wb = Workbook()
    bilgi(wb, ayar)
    hesap_plani(wb, hesaplar)
    son_hp = 4 + len(hesaplar)
    yevmiye(wb, satirlar, son_hp)
    if not sablon:
        buyuk_defter(wb, satirlar, hesaplar)
    mizan(wb, hesaplar)
    if not sablon:
        banka_ekstresi(wb, satirlar)
    if not sablon:
        kdv(wb, satirlar)
        stok(wb, sayim)
        amortisman(wb)
    gelir_tablosu(wb)
    bilanco(wb)
    if not sablon:
        takvim(wb)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = True
        ws.page_setup.orientation = "landscape"
        ws.oddFooter.center.text = "Kahve Durağı (KURGUSAL) · eğitim amaçlı · mali müşavirlik hizmeti değildir"
    wb.properties.title = "Kahve Durağı — Defter Çalışma Kitabı (kurgusal)"
    wb.properties.creator = "Muhasebe 101 lab"
    CIKTI.mkdir(exist_ok=True)
    yol = CIKTI / ("bos-sablon.xlsx" if sablon else "kahve-duragi-2026-01.xlsx")
    wb.save(yol)
    return yol


if __name__ == "__main__":
    yol = uret(sablon="--sablon" in sys.argv)
    print(yol.relative_to(DEFTER.parent))
